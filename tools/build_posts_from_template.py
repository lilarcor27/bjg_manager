"""Build the 3-6 homepage posts ⑥ and ⑦ in the layout the county made for ②.

    python3 tools/build_posts_from_template.py [제8기 위원 명단.xlsx]

With the member list, ② is rebuilt from the template with a masked
member table added as section 4 (the list itself holds personal data and
is not kept in the repository).

The county reformatted ② in Hangul (boxed title, "1." headings, "가."
sub-headings, "○" bullets, shaded header cells). This script uses that
file as the template: it keeps its styles and first paragraph (section
settings and the title box), then rebuilds the body by cloning the
template's own heading, bullet, blank and table elements, so ⑥ and ⑦
look the same and stay editable in Hangul.
"""
import copy
import os
import random
import sys
import zipfile

from lxml import etree

sys.path.insert(0, os.path.dirname(__file__))
from extract_projects import load  # noqa: E402

DIR = "docs/작성본/홈페이지 게시/"
TEMPLATE = "tools/templates/홈페이지 게시 서식(② 담당자 편집본).hwpx"  # the county's ②, kept as-is
POST2 = DIR + "② 2026년 참여예산주민위원회 구성 및 운영 현황.hwpx"
BUDGET_2026 = "docs/자료/예산 산출근거/2026 본예산 세출 산출근거(주민제안사업 노란색 표시).xlsx"
BUDGET_2027 = "docs/자료/예산 산출근거/2027 본예산(안) 주민참여예산 사업(청산면 제외).xlsx"
EUP = ["옥천읍", "동이면", "안남면", "안내면", "청성면", "청산면", "이원면", "군서면", "군북면"]
GENERIC = {"일반수용비", "일반용역비", "운영수당", "임차료", "주민참여숙원사업"}
# Given separately by the county on 2026-10-27 (not in the 2027 workbook).
CHEONGSAN_2027 = [("지역회의형", "지전리 마을안길정비공사(L=100m)", 30000), ("지역회의형", "신매리 농로확포장공사(L=100m)", 30000),
                  ("지역회의형", "덕곡리 마을안길정비공사(L=200m)", 30000), ("지역회의형", "장위리 도로정비공사(L=350m)", 35000),
                  ("지역회의형", "한곡리 도로정비공사(L=350m)", 35000), ("지역회의형", "덕지리 농로확포장공사(L=80m)", 35000),
                  ("지역회의형", "하예곡리 농로정비공사(L=140m)", 25000), ("지역회의형", "삼방리 농로및배수로정비공사(L=70m)", 25000),
                  ("지역회의형", "법화리 농로및배수로정비공사(L=100m)", 25000), ("주민자치회형", "청산 장터 야외무대 설치", 45000)]
CITIZEN_2027 = [("군민참여형", "-", "청년 AI ZONE", "성장정책과", 15000)]
SECTION = "Contents/section0.xml"
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
CENTER, LEFT = "20", "0"  # paraPrIDRef the template uses in cells

REGIONAL = [("옥천읍", "’26.9.1.", 11), ("동이면", "’26.9.10.", 18), ("안남면", "’26.8.7.", 9),
            ("안내면", "’26.8.6.", 16), ("청성면", "’26.8.14.", 37), ("청산면", "’26.8.21.", 22),
            ("이원면", "’26.9.10.", 25), ("군서면", "’26.8.26.", 12), ("군북면", "’26.9.3.", 18)]


def q(tag):
    return "{%s}%s" % (HP, tag)


def text_of(el):
    return "".join(t.text or "" for t in el.iter(q("t"))).strip()


def new_id():
    return str(random.randint(10**8, 2 * 10**9))


def drop_lineseg(el):
    for ls in list(el.iter(q("linesegarray"))):
        ls.getparent().remove(ls)


def set_lines(container_p, lines, para_pr=None):
    """Make container_p (cloned per line) hold `lines`; returns the paragraphs."""
    parent = container_p.getparent()
    out = []
    prev = container_p
    for i, line in enumerate(lines):
        p = container_p if i == 0 else copy.deepcopy(container_p)
        p.set("id", new_id())
        if para_pr is not None:
            p.set("paraPrIDRef", para_pr)
        runs = p.findall(q("run"))
        for r in runs[1:]:
            p.remove(r)
        run = runs[0]
        for child in list(run):
            run.remove(child)
        etree.SubElement(run, q("t")).text = line
        drop_lineseg(p)
        if i:
            prev.addnext(p)
        prev = p
        out.append(p)
    return out


class Template:
    def __init__(self, path, keep_body=False):
        self.path = path
        with zipfile.ZipFile(path) as z:
            raw = z.read(SECTION)
        self.decl = raw[: raw.index(b"?>") + 2]
        self.root = etree.fromstring(raw)
        paras = self.root.findall(q("p"))
        self.title_p = paras[0]
        self.heading_first = next(p for p in paras if text_of(p).startswith("1."))
        self.heading = next(p for p in paras if text_of(p).startswith("2."))
        self.sub = next(p for p in paras if text_of(p).startswith("나."))
        self.bullet = next(p for p in paras if text_of(p).startswith("○"))
        self.blank = next(p for p in paras if not text_of(p) and p.find(".//" + q("tbl")) is None
                          and p is not self.title_p)
        tables = [p for p in paras if p.find(".//" + q("tbl")) is not None and p is not self.title_p]
        self.table_p = next(p for p in tables if int(p.find(".//" + q("tbl")).get("colCnt")) >= 4)
        tbl = self.table_p.find(".//" + q("tbl"))
        rows = tbl.findall(q("tr"))
        h, b = rows[0].findall(q("tc")), rows[1].findall(q("tc"))
        self.cell = {"h": (h[0], h[1], h[-1]), "b": (b[0], b[1], b[-1])}
        self.width = int(tbl.find(q("sz")).get("width"))
        self.height = {"h": int(h[0].find(q("cellSz")).get("height")),
                       "b": int(b[0].find(q("cellSz")).get("height"))}
        self.body = paras[1:] if keep_body else []

    # ---- building blocks ------------------------------------------------
    def _clone(self, p, text=None):
        p = copy.deepcopy(p)
        p.set("id", new_id())
        if text is not None:
            set_lines(p, [text])
        drop_lineseg(p)
        self.body.append(p)

    def title(self, text):
        tbl = self.title_p.find(".//" + q("tbl"))
        cell_p = tbl.find(".//" + q("subList")).find(q("p"))
        set_lines(cell_p, [text])

    def heading_(self, text, first=False):
        self._clone(self.heading_first if first else self.heading, text)

    def sub_(self, text):
        self._clone(self.sub, "  " + text)

    def bullet_(self, text):
        self._clone(self.bullet, "    ○ " + text)

    def blank_(self):
        self._clone(self.blank)

    def table(self, cols, rows, repeat_header=False):
        """cols: [(header, weight, align)]; rows: list of lists (str or list of lines)."""
        wrap = copy.deepcopy(self.table_p)
        wrap.set("id", new_id())
        tbl = wrap.find(".//" + q("tbl"))
        tbl.set("id", new_id())
        for tr in tbl.findall(q("tr")):
            tbl.remove(tr)
        total = sum(w for _, w, _ in cols)
        widths = [self.width * w // total for _, w, _ in cols]
        widths[-1] += self.width - sum(widths)
        height = 0
        for r, values in enumerate([[c[0] for c in cols]] + rows):
            kind = "h" if r == 0 else "b"
            tr = etree.SubElement(tbl, q("tr"))
            for c, value in enumerate(values):
                first, mid, last = self.cell[kind]
                tc = copy.deepcopy(first if c == 0 else last if c == len(cols) - 1 else mid)
                tc.find(q("cellAddr")).set("colAddr", str(c))
                tc.find(q("cellAddr")).set("rowAddr", str(r))
                tc.find(q("cellSpan")).set("colSpan", "1")
                tc.find(q("cellSpan")).set("rowSpan", "1")
                tc.find(q("cellSz")).set("width", str(widths[c]))
                tc.find(q("cellSz")).set("height", str(self.height[kind]))
                if kind == "h" and repeat_header:
                    tc.set("header", "1")
                sub = tc.find(q("subList"))
                ps = sub.findall(q("p"))
                for extra in ps[1:]:
                    sub.remove(extra)
                lines = value if isinstance(value, list) else str(value).split("\n")
                align = CENTER if kind == "h" else (CENTER if cols[c][2] == "c" else LEFT)
                set_lines(ps[0], lines, align)
                tr.append(tc)
            height += self.height[kind]
        tbl.set("rowCnt", str(len(rows) + 1))
        tbl.set("colCnt", str(len(cols)))
        tbl.set("repeatHeader", "1" if repeat_header else "0")
        tbl.find(q("sz")).set("height", str(height))
        drop_lineseg(wrap)
        self.body.append(wrap)
        self.blank_()

    # ---- output ---------------------------------------------------------
    def save(self, out):
        for p in self.root.findall(q("p"))[1:]:
            self.root.remove(p)
        drop_lineseg(self.title_p)
        for p in self.body:
            self.root.append(p)
        data = self.decl + etree.tostring(self.root, encoding="UTF-8", xml_declaration=False)
        with zipfile.ZipFile(self.path) as src, zipfile.ZipFile(out + ".tmp", "w") as dst:
            for info in src.infolist():
                payload = data if info.filename == SECTION else src.read(info.filename)
                dst.writestr(info, payload, compress_type=info.compress_type)
        os.replace(out + ".tmp", out)
        print("saved", out)


REGIONAL_COLS = [("읍면", 16, "c"), ("일자", 16, "c"), ("참석(명)", 12, "c"), ("안건", 56, "l")]
REGIONAL_ROWS = [[e, d, n, f"’27년 {e} 주민참여숙원사업 대상 선정 심의"] for e, d, n in REGIONAL]


def post6():
    t = Template(TEMPLATE)
    t.title("2027년도 주민참여예산 사업 심사 과정")
    t.heading_("1. 심사 절차", first=True)
    t.table([("단계", 24, "c"), ("시기", 18, "c"), ("내용", 58, "l")], [
        ["① 주민의견 수렴", "’26.7.~8월", "‘2027년도 예산편성에 바란다’, 예산편성 방향 설문조사\n(군 홈페이지·카카오톡채널·읍면 서면)"],
        ["② 주민자치회 주민총회", "", "읍면 주민자치회가 발굴한 주민자치회형 사업을 주민총회에서 주민 투표로 선정"],
        ["③ 지역회의 심의", "’26.8.6.~9.10.", "9개 읍면 지역회의에서 주민참여숙원사업 대상 선정 심의"],
        ["④ 소관부서 검토", "", "접수된 주민제안사업을 소관부서가 검토(제안자와 소통하여 사업 구체화)"],
        ["⑤ 분과위원회 심의", "’26.10.12.~10.16.", "3개 분과위원회에서 사업설명 청취, 질의답변, 의견 제시 후\n대상사업 조정 및 우선순위 최종결정"],
        ["⑥ 예산안 편성", "", "예산안 편성 및 주민의견서 의회 제출"],
    ])
    t.heading_("2. 지역회의 심의")
    t.table(REGIONAL_COLS, REGIONAL_ROWS)
    t.heading_("3. 분과위원회 심의")
    t.table([("분과", 13, "c"), ("일시", 17, "c"), ("참석/\n재적", 9, "c"), ("계속\n(건)", 8, "c"),
             ("신규\n(건)", 8, "c"), ("’27 요구액\n(천원)", 15, "c"), ("’26 예산액\n(천원)", 15, "c"),
             ("증감\n(천원)", 15, "c")], [
        ["자치행정분과", "’26.10.12.(월)\n20:00~21:00", "6/11", "90", "10", "5,708,626", "4,919,996", "788,630"],
        ["농산개발분과", "’26.10.13.(화)\n10:00~11:20", "8/11", "98", "5", "8,992,133", "9,985,099", "△992,966"],
        ["복지문화분과", "’26.10.16.(금)\n15:00~16:30", "9/11", "175", "12", "5,618,733", "5,242,467", "376,266"],
    ])
    t.heading_("4. 분과위원회 예산편성 방향 의견")
    t.bullet_("(자치행정) 보조사업 집행 단가 현실화 필요, 순세계잉여금 등을 활용한 체육진흥기금의 계획적 적립")
    t.bullet_("(농산개발) 특정 축종에 집중된 지원을 다른 축종으로 확대, 성과 부진 사업의 추진상황 지속 점검")
    t.bullet_("(복지문화) 고령화 여건을 고려한 주민 체감형 편성, 경로당 등 이용 실태에 맞는 지원")
    t.save(DIR + "⑥ 2027년도 주민참여예산 사업 심사 과정.hwpx")


def projects_2027():
    """2027 projects selected in 2026: the county's 2027 draft-budget export
    (level 1 ◎ / level 2 ○ lines), plus 청산면 and the citizen-type project."""
    import openpyxl

    ws = openpyxl.load_workbook(BUDGET_2027, data_only=True)["Sheet0"]
    head = [c.value for c in ws[1]]
    i_dept, i_biz = head.index("부서명"), head.index("세부사업명")
    i_lvl, i_name, i_amt = head.index("조정레벨"), head.index("조정산출근거"), head.index("경정조정액")
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] == "2027"]
    groups = {}
    for r in rows:
        groups.setdefault((r[i_dept], r[i_biz]), []).append((r[i_lvl], r[i_name].strip(), r[i_amt]))
    found = []
    for (dept, biz), lines in groups.items():
        typ = "지역회의형" if "숙원" in biz else "주민자치회형"
        for k, (lvl, name, amt) in enumerate(lines):
            if lvl != "1":
                continue
            kids = []
            for lvl2, name2, amt2 in lines[k + 1:]:
                if lvl2 == "1":
                    break
                kids.append((name2.lstrip("○").strip(), amt2))
            head_name = name.lstrip("◎").strip()
            picks = kids if kids and head_name in GENERIC else [(head_name, amt)]
            for nm, am in picks:
                eup = next((e for e in EUP if e in dept or f"({e})" in nm), "")
                found.append({"typ": typ, "eup": eup, "name": nm, "dept": dept, "amt": am})
    found += [{"typ": t, "eup": "청산면", "name": n, "dept": "청산면", "amt": a} for t, n, a in CHEONGSAN_2027]
    found += [{"typ": t, "eup": e, "name": n, "dept": d, "amt": a} for t, e, n, d, a in CITIZEN_2027]
    order = {"지역회의형": 0, "주민자치회형": 1, "군민참여형": 2}
    found.sort(key=lambda x: (order[x["typ"]], EUP.index(x["eup"]) if x["eup"] in EUP else 9, x["dept"] != "행정과"))
    return found


def post7():
    """⑦ 2027 selection results: projects chosen through this year's process."""
    items = projects_2027()
    total = sum(e["amt"] for e in items)
    by_type = {}
    for e in items:
        by_type.setdefault(e["typ"], [0, 0])
        by_type[e["typ"]][0] += 1
        by_type[e["typ"]][1] += e["amt"]
    labels = {"지역회의형": "지역회의형(주민참여숙원사업)", "주민자치회형": "주민자치회형(주민자치회 참여예산)",
              "군민참여형": "군민참여형"}
    t = Template(TEMPLATE)
    t.title("2027년도 주민참여예산 공모사업 선정 결과")
    t.heading_("1. 선정 개요", first=True)
    t.bullet_(f"2026년 주민참여예산 공모·심사를 거쳐 선정된 2027년도 사업: {len(items)}건, {total:,}천원")
    for typ, (n, amt) in by_type.items():
        t._clone(t.bullet, f"      - {labels[typ]} {n}건, {amt:,}천원")
    t.bullet_("2027년도 본예산(안) 기준이며, 군의회 심의 결과에 따라 변경될 수 있습니다.")
    t.blank_()
    t.heading_("2. 사업별 내역")
    rows = [["합계", "", "", "", "", f"{total:,}"]]
    rows += [[i, e["typ"], e["eup"] or "-", e["name"], e["dept"], f"{e['amt']:,}"] for i, e in enumerate(items, 1)]
    t.table([("연번", 7, "c"), ("유형", 14, "c"), ("읍면", 10, "c"), ("사업명", 44, "l"),
             ("소관부서", 12, "c"), ("예산액(천원)", 13, "c")], rows, repeat_header=True)
    t.save(DIR + "⑦ 2027년도 주민참여예산 공모사업 선정 결과.hwpx")


def post7_2026():
    items, _issues, _total = load(BUDGET_2026, 2026)
    total = sum(e["amt"] for e in items)
    local = sum(e["amt"] for e in items if e["typ"] == "지역회의형")
    t = Template(TEMPLATE)
    t.title("2026년도 주민참여예산 주민제안사업 선정 결과")
    t.heading_("1. 선정 개요", first=True)
    t.bullet_(f"2026년도 본예산에 편성된 주민제안사업: {len(items)}건, {total:,}천원")
    t.bullet_(f"지역회의형(주민참여숙원사업) {local:,}천원, 주민자치회형(주민자치회 참여예산) {total - local:,}천원")
    t.blank_()
    t.heading_("2. 사업별 내역")
    rows = [["합계", "", "", "", "", f"{total:,}"]]
    rows += [[i, e["typ"], e["eup"], e["name"], e["dept"], f"{e['amt']:,}"] for i, e in enumerate(items, 1)]
    t.table([("연번", 7, "c"), ("유형", 14, "c"), ("읍면", 10, "c"), ("사업명", 44, "l"),
             ("소관부서", 12, "c"), ("예산액(천원)", 13, "c")], rows, repeat_header=True)
    t.save(DIR + "⑦ 2026년도 주민참여예산 주민제안사업 선정 결과.hwpx")


def members(path):
    """Public member list: names masked to the surname, no birth date,
    address, phone, job or attendance."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True)["전체"]
    out = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not isinstance(row[0], int):
            continue
        name = str(row[3]).strip()
        out.append({"subc": row[1], "role": row[2] or "위원", "name": name[0] + "○" * (len(name) - 1),
                    "sex": row[5], "note": row[9]})
    return out


def update2(member_xlsx):
    rows = members(member_xlsx)
    n = len(rows)
    youth = sum("청년" in (r["note"] or "") for r in rows)
    local = sum("지역회의" in (r["note"] or "") for r in rows)
    men = sum(r["sex"] == "남" for r in rows)
    t = Template(TEMPLATE, keep_body=True)
    t.heading_("4. 제8기 참여예산주민위원회 위원 명단")
    t.bullet_(f"총 {n}명(남 {men}명, 여 {n - men}명), 청년위원 {youth}명({youth * 100 // n}%), "
              f"지역회의 추천 {local}명, 임기 ’25.10.10.~’27.10.9.")
    t.blank_()
    t.table([("연번", 8, "c"), ("분과", 18, "c"), ("직책", 14, "c"), ("성명", 14, "c"),
             ("성별", 10, "c"), ("구분", 36, "c")],
            [[i, r["subc"], r["role"], r["name"], r["sex"], r["note"]] for i, r in enumerate(rows, 1)],
            repeat_header=True)
    t.bullet_("개인정보 보호를 위해 성명 일부를 가리고, 생년월일·주소·연락처는 공개하지 않습니다.")
    t.save(POST2)


if __name__ == "__main__":
    post6()
    post7()
    if len(sys.argv) > 1:
        update2(sys.argv[1])
