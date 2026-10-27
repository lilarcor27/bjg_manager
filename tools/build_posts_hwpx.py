"""Build the 3-6 homepage posts (②, ⑥, ⑦) as editable HWPX files.

    python3 tools/build_posts_hwpx.py

Same content as the PDFs in docs/작성본/홈페이지 게시/, so a successor can
edit them in Hangul. Uses python-hwpx (pip install python-hwpx). The 2026
project list for ⑦ comes from the yellow-marked budget workbook.
"""
import os
import sys

import openpyxl
from hwpx import HwpxDocument

sys.path.insert(0, os.path.dirname(__file__))

OUT = "docs/작성본/홈페이지 게시/"
FONT = "맑은 고딕"
HEAD_FILL = "#E8EDF5"
TOTAL_FILL = "#F4F4F4"
BUDGET_2026 = "docs/자료/예산 산출근거/2026 본예산 세출 산출근거(주민제안사업 노란색 표시).xlsx"

REGIONAL = [("옥천읍", "’26.9.1.", 11), ("동이면", "’26.9.10.", 18), ("안남면", "’26.8.7.", 9),
            ("안내면", "’26.8.6.", 16), ("청성면", "’26.8.14.", 37), ("청산면", "’26.8.21.", 22),
            ("이원면", "’26.9.10.", 25), ("군서면", "’26.8.26.", 12), ("군북면", "’26.9.3.", 18)]
SUBS = [("자치행정분과", "’26.10.12.(월) 20:00~21:00", "6/11", "90", "10", "5,708,626", "4,919,996", "788,630"),
        ("농산개발분과", "’26.10.13.(화) 10:00~11:20", "8/11", "98", "5", "8,992,133", "9,985,099", "△992,966"),
        ("복지문화분과", "’26.10.16.(금) 15:00~16:30", "9/11", "175", "12", "5,618,733", "5,242,467", "376,266")]


class Post:
    def __init__(self, title):
        self.doc = HwpxDocument.new()
        self.doc.page.set_margins(left=5670, right=5670, top=5670, bottom=4252)  # 20/20/20/15 mm
        self.body = self.doc.styles.ensure_run(font=FONT, size=10)
        self.bold = self.doc.styles.ensure_run(font=FONT, size=10, bold=True)
        self.h2 = self.doc.styles.ensure_run(font=FONT, size=12, bold=True)
        self.note = self.doc.styles.ensure_run(font=FONT, size=10, color="#444444")
        self.para(title, self.doc.styles.ensure_run(font=FONT, size=17, bold=True), align="CENTER")
        self.para("옥천군 기획예산담당관 · 2026. 10.", self.note, align="CENTER", after=8)

    def para(self, text, style=None, align=None, after=None, before=None):
        self.doc.add_paragraph(text, char_pr_id_ref=style or self.body)
        idx = len(self.doc.paragraphs) - 1
        fmt = {}
        if align:
            fmt["alignment"] = align
        if after is not None:
            fmt["spacing_after_pt"] = after
        if before is not None:
            fmt["spacing_before_pt"] = before
        if fmt:
            self.doc.styles.apply_paragraph_format(paragraph_index=idx, **fmt)

    def heading(self, text):
        self.para(text, self.h2, before=8, after=3)

    def table(self, rows, widths, header_rows=1, merges=(), shade_rows=()):
        t = self.doc.add_table(len(rows), len(rows[0]), char_pr_id_ref=self.body)
        t.set_column_widths(widths)
        for (r1, c1, r2, c2) in merges:
            t.merge_cells(r1, c1, r2, c2)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                if val is None:
                    continue
                t.set_cell_text(r, c, str(val), split_paragraphs=True)
        for r in range(header_rows):
            for c in range(len(rows[0])):
                if rows[r][c] is not None:
                    t.set_cell_shading(r, c, HEAD_FILL)
        for r in shade_rows:
            for c in range(len(rows[0])):
                if rows[r][c] is not None:
                    t.set_cell_shading(r, c, TOTAL_FILL)
        self.para("")
        return t

    def save(self, name):
        path = OUT + name
        self.doc.save_to_path(path)
        print("saved", path)


def regional_rows():
    return [["읍면", "일자", "참석(명)", "안건"]] + [
        [e, d, n, f"’27년 {e} 주민참여숙원사업 대상 선정 심의"] for e, d, n in REGIONAL]


def post2():
    p = Post("2026년 참여예산주민위원회 구성 및 운영 현황")
    p.heading("1. 참여예산주민위원회")
    p.table([
        ["구분", "내용"],
        ["설치 근거", "「옥천군 주민참여예산제 운영조례」 제8조(설치), 제16조(분과위원회)"],
        ["구성", "제8기 위원 33명 (자치행정·농산개발·복지문화 3개 분과, 분과별 11명)"],
        ["임기", "2년(’25.10.~’27.10.), 1회 연임 가능"],
        ["위촉 방법", "공개모집 및 주민참여예산 지역회의 추천 병행"],
        ["선정 기준", "지역(읍면), 성별, 사회적약자(노인, 장애인) 등 고려, 신규위원 우선 선정\n"
                    "청년의 의사반영 및 참여보장을 위해 청년위원 20% 이상 선정"],
        ["주요 기능", "예산편성에 대한 주민의견 수렴·집약, 중점 투자사업 우선순위 결정 및 의견 제출,\n"
                    "주민 대상 예산 홍보, 보고회·토론회 개최 등(조례 제9조)"],
    ], [22, 78])
    p.heading("2. 주민참여예산 지역회의")
    p.table([
        ["구분", "내용"],
        ["구성", "9개 읍면, 읍면별 30명 이내\n(참여 희망 주민 및 주민자치회 등 읍면 단위 단체 회원 중 읍면장 위촉)"],
        ["임기", "2년(참여예산주민위원회 임기와 동일)"],
        ["주요 기능", "예산편성 관련 중점 투자분야 주민의견 수렴, 주민참여숙원사업 결정,\n참여예산주민위원회 위원 추천"],
    ], [22, 78])
    p.heading("3. 2026년 운영 현황")
    p.para("가. 회의 개최: 총 12회 (분과위원회 3회, 지역회의 9회)", p.bold, after=2)
    p.table([["분과", "일시", "참석/재적", "안건"]] + [
        [s[0], s[1], s[2], "’27년 당초예산 주민참여예산사업 사업별 의견제출 및 우선순위 결정"] for s in SUBS],
        [16, 30, 12, 42])
    p.table(regional_rows(), [16, 16, 12, 56])
    p.para("나. 교육", p.bold, after=2)
    p.para("○ 2026 주민참여예산학교: ’26.9.29.(화) 16:00~17:30, 참여예산주민위원 17명\n"
           "   (특강 ‘옥천군 주민참여예산 협력적 소통으로 참여하기’)")
    p.para("○ 읍면 주민참여예산 홍보 및 교육: ’26.7.8.~8.28., 9개 읍면 35회\n"
           "   (주민자치회, 이장협의회, 새마을단체, 기관단체협의회 등)")
    p.save("② 2026년 참여예산주민위원회 구성 및 운영 현황.hwpx")


def post6():
    p = Post("2027년도 주민참여예산 사업 심사 과정")
    p.heading("1. 심사 절차")
    p.table([
        ["단계", "시기", "내용"],
        ["① 주민의견 수렴", "’26.7.~8월", "‘2027년도 예산편성에 바란다’, 예산편성 방향 설문조사\n(군 홈페이지·카카오톡채널·읍면 서면)"],
        ["② 주민자치회 주민총회", "", "읍면 주민자치회가 발굴한 주민자치회형 사업을 주민총회에서 주민 투표로 선정"],
        ["③ 지역회의 심의", "’26.8.6.~9.10.", "9개 읍면 지역회의에서 주민참여숙원사업 대상 선정 심의"],
        ["④ 소관부서 검토", "", "접수된 주민제안사업을 소관부서가 검토(제안자와 소통하여 사업 구체화)"],
        ["⑤ 분과위원회 심의", "’26.10.12.~10.16.", "3개 분과위원회에서 사업설명 청취, 질의답변, 의견 제시 후\n대상사업 조정 및 우선순위 최종결정"],
        ["⑥ 예산안 편성", "", "예산안 편성 및 주민의견서 의회 제출"],
    ], [24, 18, 58])
    p.heading("2. 지역회의 심의")
    p.table(regional_rows(), [16, 16, 12, 56])
    p.heading("3. 분과위원회 심의")
    rows = [["분과", "일시", "참석/재적", "대상 사업(건)", None, "예산(천원)", None, None],
            [None, None, None, "계속", "신규", "’27 요구액", "’26 예산액", "증감"]]
    rows += [list(s) for s in SUBS]
    p.table(rows, [13, 25, 9, 7, 7, 13, 13, 13], header_rows=2,
            merges=[(0, 0, 1, 0), (0, 1, 1, 1), (0, 2, 1, 2), (0, 3, 0, 4), (0, 5, 0, 7)])
    p.heading("4. 분과위원회 예산편성 방향 의견")
    p.para("○ (자치행정) 보조사업 집행 단가 현실화 필요, 순세계잉여금 등을 활용한 체육진흥기금의 계획적 적립")
    p.para("○ (농산개발) 특정 축종에 집중된 지원을 다른 축종으로 확대, 성과 부진 사업의 추진상황 지속 점검")
    p.para("○ (복지문화) 고령화 여건을 고려한 주민 체감형 편성, 경로당 등 이용 실태에 맞는 지원")
    p.save("⑥ 2027년도 주민참여예산 사업 심사 과정.hwpx")


def projects_2026():
    """Yellow column-I rows of the 2026 budget, grouped like the 1-1 evidence list."""
    from extract_projects import load  # noqa: E402
    entries, _issues, _total = load(BUDGET_2026, 2026)
    return entries


def post7():
    items = projects_2026()
    total = sum(e["amt"] for e in items)
    local = sum(e["amt"] for e in items if e["typ"] == "지역회의형")
    p = Post("2026년도 주민참여예산 주민제안사업 선정 결과")
    p.para(f"2026년도 본예산에 편성된 주민제안사업은 {len(items)}건, {total:,}천원입니다. "
           f"(지역회의형 {local:,}천원, 주민자치회형 {total - local:,}천원)", after=4)
    rows = [["연번", "유형", "읍면", "사업명", "소관부서", "예산액(천원)"],
            ["합계", None, None, None, None, f"{total:,}"]]
    rows += [[i, e["typ"], e["eup"], e["name"], e["dept"], f"{e['amt']:,}"] for i, e in enumerate(items, 1)]
    p.table(rows, [7, 14, 10, 45, 11, 13], merges=[(1, 0, 1, 4)], shade_rows=[1])
    p.save("⑦ 2026년도 주민참여예산 주민제안사업 선정 결과.hwpx")


if __name__ == "__main__":
    post2()
    post6()
    post7()
